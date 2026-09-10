from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
LOGO_PATH = ROOT / "public" / "logo" / "android-chrome-512x512.png"
DOCX_PATH = DOCS_DIR / "deepfense-website-description-en.docx"
MD_PATH = DOCS_DIR / "deepfense-website-description-en.md"

BLUE = "1D6FE8"
DARK = "0F172A"
SLATE = "475569"
LIGHT_BLUE = "EAF3FF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = DARK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(document: Document, text: str, style: str | None = None, bold_prefix: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.08

    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        prefix.font.color.rgb = RGBColor.from_string(DARK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        rest.font.color.rgb = RGBColor.from_string(DARK)
    else:
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor.from_string(DARK)

    return paragraph


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DARK)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DARK)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    heading.paragraph_format.space_after = Pt(5)
    for run in heading.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else DARK)
        run.bold = True


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_shading(header_cells[index], BLUE)
        set_cell_text(header_cells[index], header, bold=True, color="FFFFFF")
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if index == 0:
                set_cell_shading(cells[index], LIGHT_BLUE)

    document.add_paragraph()


def build_markdown() -> str:
    return """# DEEPFENSE.ONLINE Website Description

**Team:** 2HAND  
**Members:** Ho Xuan Nguyen (25NS039), Nguyen Nhat Huy (25NS020), Vo Phan Phap (25NS045), Tran Minh Nhat (25NS040)  
**University:** Vietnam-Korea University of Information and Communication Technology (VKU), Class 25NS  
**Website:** https://deepfense.online  
**Product type:** Gamified deepfake awareness, AI scam defense, and digital-safety training platform

## 1. Executive Overview

DEEPFENSE.ONLINE is a web platform designed to help everyday users recognize deepfakes, AI voice scams, synthetic media manipulation, and social-engineering scenarios before they make unsafe decisions. The project does not position detection as a single final answer. Instead, it combines education, practice, risk review, AI guidance, emergency response, and progress recognition into one user journey.

The core idea is simple: deepfake defense is not only a technical problem. It is also a human reaction problem. A realistic face, a familiar voice, an urgent request, or a convincing image can push users to act before they verify. Deepfense trains users to slow down, observe suspicious cues, verify through safer channels, and respond correctly when something has already gone wrong.

In short, Deepfense aims to turn awareness into a practical reflex:

**Stop -> Observe -> Verify -> Respond.**

## 2. Problem Statement

Deepfake and AI scam risks are becoming more personal, faster, and harder to judge by sight alone. Attackers no longer need a perfect fake to create damage. They only need enough realism, a believable story, and pressure that makes the victim react quickly.

Deepfake threats can affect:

- Trust: users may believe a fake face, fake voice, or fake message because it appears to come from someone familiar.
- Finance: scams may request urgent money transfers, investment decisions, wallet connections, or account access.
- Identity: public photos, videos, and voice clips can be reused to impersonate a real person.
- Reputation and mental health: manipulated images or videos can be used for harassment, blackmail, humiliation, or false accusations.
- Community safety: fake content can spread quickly through group chats, social platforms, and livestreams.

Many existing tools focus mainly on whether a media file is fake. Deepfense expands the response: it teaches users why a situation is risky, lets them practice recognition, gives them a risk-scoring workflow, and provides guidance for what to do next.

## 3. Target Users

Deepfense is built for users who may not be cybersecurity specialists but still face AI-driven threats in daily life.

- Students and young social-media users who need to understand synthetic images, fake video calls, impersonation, and unsafe sharing behavior.
- Families and elderly users who are vulnerable to urgent calls, family impersonation, romance scams, and fake authority messages.
- Teachers, schools, and universities that need structured digital-safety training.
- Small organizations and businesses that need awareness training for staff before AI scams become operational incidents.
- Community users who want a clear place to learn, practice, ask questions, scan suspicious situations, and receive first-response guidance.

## 4. Product Philosophy

Deepfense is built around a practical learning philosophy: a user should not only read about deepfake risks, but also experience realistic decisions in a safe environment.

The website uses gamified learning, scenario simulation, instant feedback, and visible progress. This helps users build habits instead of memorizing isolated warning signs. The platform also avoids promising perfect detection. Risk scores and AI responses are treated as decision-support signals that should be combined with source verification, context checks, and safe follow-up actions.

The character AN can act as a guidance layer inside the learning story. AN is not the product itself, but a companion that helps connect lessons, challenges, tools, and crisis response into a coherent journey.

## 5. Website Journey

The website experience is designed as a complete defense path:

1. Home: the user understands what Deepfense is and where to begin.
2. Academy: the user learns concepts, warning signs, verification habits, and safe response patterns.
3. Challenge: the user practices recognition through deepfake detective tasks and simulated scam pressure.
4. Tools / Risk Scan: the user reviews suspicious behavior or media through structured risk workflows.
5. AI Assistant: the user asks questions, pastes suspicious content, and receives guided safety advice.
6. Crisis Hub: the user receives first-aid steps if a scam, harassment, blackmail, or data exposure has already happened.
7. Profile / Certificate: the user can track progress, unlock certificate-related milestones, and demonstrate training completion.

This journey makes the website more than an information page. It becomes a training loop: learn, practice, scan, ask, respond, and improve.

## 6. Main Website Sections and Features

### 6.1 Home

The Home page introduces Deepfense as a deepfake awareness and AI scam defense training platform. It gives users clear entry points into the main actions: learning signals, practicing real scenarios, using risk tools, and earning recognition through certificates or DPF-related progress.

The Home page also frames the product around community safety. It is not only built for experts. It is designed so that a user can quickly understand the risk and choose a next step without feeling overwhelmed by a complex security dashboard.

### 6.2 Academy

Academy is the structured learning area. It helps users understand what deepfakes are, how synthetic media differs from normal edited content, why AI voice and video impersonation are dangerous, and why verification habits matter.

Key Academy capabilities include:

- Foundation course content for deepfake concepts, deepvoice, synthetic media, edited media, social engineering, and verification.
- Lesson modules with quizzes and interactive learning blocks.
- Progress saving through user authentication.
- Certificate and DPF-related milestones after required learning and evaluation tasks are completed.
- A training tone focused on practical judgment, not fear.

Academy is important because scanning tools alone cannot protect users in every situation. Users need mental models that help them pause, check the source, question urgency, and verify through independent channels.

### 6.3 Challenge

The Challenge area turns knowledge into active practice. The Deepfake Detective Challenge uses a multi-level training flow where users inspect suspicious media or scenarios, choose answers, receive feedback, and review their result.

Key Challenge capabilities include:

- A 10-level detection and awareness flow.
- Scoring and result feedback.
- Post-game analysis that helps users understand what they missed.
- Survey and sharing options to support awareness and product improvement.
- A scenario simulator that exposes users to pressure tactics, urgent requests, and manipulation patterns.

This section is designed to make mistakes useful. If a user chooses the wrong answer, the system can explain which cue was missed and how to respond better next time.

### 6.4 Tools and Risk Scan

The Tools page is the operational part of the website. It includes scan and forensics workflows, a protective shield concept, a knowledge base, and the Crisis Hub.

Risk Scan combines two layers:

- Behavioral risk review: the user answers scenario-based questions about urgency, identity, request type, money, links, OTPs, or unusual behavior. The system estimates risk and recommends safer actions.
- Media and forensics review: public AI scanning is currently locked while datasets, model training, and benchmarks are prepared. Future review workflows should present evidence, provenance, limitations, and risk-score display only after validation.

The page also presents rPPG/liveness concepts, inspired by the idea that real human faces contain biological signals such as subtle blood-flow changes that synthetic faces may fail to reproduce. This communicates that deepfake defense can combine technical signals and human verification behavior.

### 6.5 Protective Shield

Protective Shield is a preventive layer. It is based on a Fawkes-style concept: injecting invisible adversarial noise into personal photos to reduce the chance that AI models can use them for unauthorized face recognition or training.

The feature supports the broader Deepfense message: users should not wait until deepfake abuse happens. They can reduce exposure, protect personal media, and understand that public images and voice clips can become raw material for impersonation.

### 6.6 Knowledge and Law

Knowledge and Law gives users structured references about AI, deepfake technology, digital rights, detection methods, and legal or ethical context. It helps users understand the issue beyond a single scam case.

Topics include:

- Deepfake and AI-generated media basics.
- Technical ideas such as rPPG, C2PA, and spectrogram analysis.
- AI law, labeling, watermarking, and human rights in the digital age.
- Future trends and the need for proactive defense.

This section supports judges, educators, and learners who want to see that the website is grounded in education and responsible use, not only entertainment.

### 6.7 AI Assistant

The AI Assistant is available as a chat-based support layer. It can help users ask questions about deepfake, online scams, suspicious links, media risks, and emergency response steps.

Its intended role is a safety coach:

- Explain suspicious signs in simple language.
- Help analyze scam scripts and manipulation tactics.
- Guide users through safer verification steps.
- Provide first-response suggestions when a user is confused or under pressure.

The assistant should be understood as guidance, not a replacement for professional legal, banking, platform, or law-enforcement support.

### 6.8 Crisis Hub

Crisis Hub focuses on what happens after a user is harmed or feels at immediate risk. This is important because deepfake incidents often create panic, shame, and confusion. Users need calm, step-by-step actions.

Core Crisis Hub goals include:

- Psychological first aid: help the user stop, breathe, and avoid making the situation worse.
- Account and data protection: change passwords, enable two-factor authentication, revoke sessions, and protect payment accounts.
- Evidence preservation: save screenshots, URLs, transaction records, account names, timestamps, and chat logs.
- Reporting guidance: direct users toward appropriate reporting channels and support workflows.
- Harm reduction: avoid forwarding sensitive content, avoid negotiating under pressure, and seek trusted help.

Crisis Hub does not replace official authorities. It provides the first response path that helps users act correctly before they escalate the case.

### 6.9 Profile, Certificate, and Admin

The Profile area helps authenticated users manage their learner identity and progress. The certificate verification route supports the idea that training completion can be shown and checked.

The Admin route is protected and supports operational management for roles such as admin or editor. This is relevant for future institutional use, where schools or organizations may need to manage users, training records, challenges, and learning data.

## 7. Feature Summary

| Feature | Purpose | User Outcome |
| --- | --- | --- |
| Home | Introduce the platform and guide the user to the right starting point. | The user knows what Deepfense does and where to begin. |
| Academy | Teach deepfake concepts, verification habits, and safe response. | The user builds knowledge before facing real threats. |
| Challenge | Convert knowledge into practice through levels, scoring, and feedback. | The user improves recognition and decision-making reflexes. |
| Risk Scan | Review suspicious behavior or media with risk-oriented workflows. | The user receives a risk signal and recommended next actions. |
| Protective Shield | Reduce exposure of personal photos to unauthorized AI training. | The user learns prevention, not only detection. |
| Knowledge and Law | Explain AI, deepfake technology, detection methods, rights, and trends. | The user understands the wider digital-safety context. |
| AI Assistant | Provide conversational guidance for scams, deepfake questions, and safety steps. | The user can ask for help when confused or under pressure. |
| Crisis Hub | Give emergency first-response guidance after a potential incident. | The user knows what to do next and how to preserve evidence. |
| Profile and Certificate | Track learning progress and recognize completion. | The user can show training progress or completion. |

## 8. Technical Stack

Deepfense is implemented as a modern React web application.

| Layer | Technologies |
| --- | --- |
| Frontend | React 19, Vite, TypeScript |
| Styling and UI | Tailwind CSS, Lucide Icons |
| Motion and 3D | GSAP, Lenis, Three.js, React Three Fiber |
| Authentication and Data | Firebase Authentication, Firebase Firestore |
| AI and Serverless | Google Gemini API, Vercel Serverless Functions |
| Hosting | Vercel |

The stack supports an interactive learning experience with animated pages, scenario-based modules, authenticated progress, AI-assisted guidance, and deployable web infrastructure.

## 9. Product Value

Deepfense creates value by connecting several layers that are usually separated:

- Education: users learn the core concepts behind deepfake and AI scam risk.
- Practice: users train through challenges instead of only reading static content.
- Decision support: risk scan and assistant workflows help users evaluate suspicious situations.
- Emergency response: Crisis Hub gives practical next steps when damage may already be happening.
- Recognition: certificate and DPF-related progress motivate users and support institutional training.

The strongest value proposition is that Deepfense does not only ask "Is this fake?" It asks the more useful question: "What should the user do safely now?"

## 10. Business and Deployment Potential

Deepfense can grow through a community-first model while still supporting sustainable monetization.

Potential models include:

- B2B or B2E training packages for schools, universities, small businesses, and organizations.
- Premium learning modules for advanced detection, reporting, and organizational defense.
- Certificate-based training completion for institutions.
- Sponsored digital-safety campaigns with education partners, community groups, or cybersecurity organizations.
- Custom dashboards for schools or organizations to track training participation and completion.

The basic public-facing knowledge layer should remain accessible because deepfake defense has community value. Revenue should come from structured training, institutional reporting, advanced modules, and support services rather than blocking essential safety information.

## 11. Future Development Direction

The current website can become the foundation for a larger Deepfense ecosystem.

Planned or recommended directions include:

- Mobile app: provide faster access during real calls, chats, and urgent scam situations.
- Organization dashboard: allow schools and companies to manage training cohorts, completion rates, and risk awareness.
- Community reporting: collect anonymized scam patterns and deepfake scenarios to improve education materials.
- Better AI Assistant flow: improve triage, safer response templates, and source-verification guidance.
- Dataset and challenge expansion: add more realistic cases across finance, family, romance, authority, and social-media contexts.
- Browser extension or real-time protection layer: help users identify risky media and links closer to where attacks happen.
- Privacy and data governance improvements: minimize sensitive uploads, clarify retention, and strengthen user trust.

## 12. Conclusion

DEEPFENSE.ONLINE is not only a deepfake detector. It is a training platform for digital defense behavior. Its main contribution is the combination of learning, practice, scan workflows, AI guidance, crisis response, and progress recognition in one website.

As deepfakes become easier to create, the safest user will not be the person who guesses fastest. The safest user will be the person who knows how to pause, observe, verify, and respond. Deepfense is built to train that reflex.
"""


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Aptos Display"


def add_cover(document: Document) -> None:
    if LOGO_PATH.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(LOGO_PATH), width=Inches(1.05))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("DEEPFENSE.ONLINE")
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(28)
    title_run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Website Description and Product Overview")
    subtitle_run.font.name = "Aptos Display"
    subtitle_run.font.size = Pt(15)
    subtitle_run.font.color.rgb = RGBColor.from_string(DARK)

    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline_run = tagline.add_run("Gamified Deepfake Awareness, AI Scam Defense, and Digital-Safety Training")
    tagline_run.italic = True
    tagline_run.font.name = "Aptos"
    tagline_run.font.size = Pt(10.5)
    tagline_run.font.color.rgb = RGBColor.from_string(SLATE)

    document.add_paragraph()
    info = [
        ["Team", "2HAND"],
        ["Members", "Ho Xuan Nguyen (25NS039), Nguyen Nhat Huy (25NS020), Vo Phan Phap (25NS045), Tran Minh Nhat (25NS040)"],
        ["University", "Vietnam-Korea University of Information and Communication Technology (VKU), Class 25NS"],
        ["Website", "https://deepfense.online"],
        ["Audience", "Final-round judges and product reviewers"],
    ]
    add_table(document, ["Field", "Description"], info)
    document.add_page_break()


def add_markdown_to_docx(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    in_table = False
    table_headers: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal in_table, table_headers, table_rows
        if in_table and table_headers:
            add_table(document, table_headers, table_rows)
        in_table = False
        table_headers = []
        table_rows = []

    for raw_line in lines:
        line = raw_line.rstrip()
        if line.startswith("# DEEPFENSE.ONLINE"):
            continue

        if line.startswith("| ") and line.endswith(" |"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(set(cell) <= {"-", ":", " "} for cell in cells):
                continue
            if not in_table:
                in_table = True
                table_headers = cells
                table_rows = []
            else:
                table_rows.append(cells)
            continue
        flush_table()

        if not line:
            continue
        if line.startswith("## "):
            add_heading(document, line[3:], level=1)
        elif line.startswith("### "):
            add_heading(document, line[4:], level=2)
        elif line.startswith("- "):
            add_bullet(document, line[2:])
        elif line[:3].strip(".").isdigit() and ". " in line[:5]:
            add_number(document, line.split(". ", 1)[1])
        elif line.startswith("**") and line.endswith("**") and line.count("**") == 2:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(line.strip("*"))
            run.bold = True
            run.font.name = "Aptos Display"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif line.startswith("**") and ":**" in line:
            label, value = line.split(":**", 1)
            add_paragraph(document, f"{label.strip('*')}: {value.strip()}", bold_prefix=f"{label.strip('*')}:")
        else:
            add_paragraph(document, line)

    flush_table()


def add_footer(document: Document) -> None:
    section = document.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("DEEPFENSE.ONLINE | 2HAND | Website Description")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(SLATE)


def main() -> None:
    DOCS_DIR.mkdir(exist_ok=True)
    markdown = build_markdown()
    MD_PATH.write_text(markdown, encoding="utf-8")

    document = Document()
    configure_document(document)
    add_cover(document)
    add_markdown_to_docx(document, markdown)
    add_footer(document)
    document.core_properties.title = "DEEPFENSE.ONLINE Website Description"
    document.core_properties.subject = "Product overview for final-round judges"
    document.core_properties.author = "2HAND"
    document.save(DOCX_PATH)

    print(f"Created {DOCX_PATH}")
    print(f"Created {MD_PATH}")


if __name__ == "__main__":
    main()
